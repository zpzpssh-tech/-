# 영업/원본조사/list_*.json 을 합쳐 영업/B2B_영업리스트.xlsx 를 만듭니다. 실행: python3 scripts/build_sales_list.py
import json, glob, os, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

SCR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "영업", "원본조사")
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "영업")
rows = []
for f in sorted(glob.glob(f"{SCR}/list_*.json")):
    try:
        data = json.load(open(f, encoding="utf-8"))
    except Exception as e:
        print("skip", f, e); continue
    for d in data:
        rows.append(d)

# dedupe by homepage domain, then by normalized name; merge fields
import re
from urllib.parse import urlparse
EXCLUDE={"준보코리아","루아샵","기프트조아","도매토피아","VBTRADE"}
def dom(u):
    try: return urlparse(u if u.startswith("http") else "http://"+u).netloc.lower().replace("www.","").replace("m.","") if u else ""
    except: return ""
def nm(n):
    n=re.sub(r"\(.*?\)","",n); return re.sub(r"\s|\(주\)|주식회사|㈜","",n)
merged={}; order=[]
for r in rows:
    name=r.get("업체명","")
    if any(e in name for e in EXCLUDE): continue
    key=dom(r.get("홈페이지","")) or nm(name)
    alt=nm(name)
    k=None
    for kk in (key,alt):
        if kk in merged: k=kk; break
    if k is None:
        merged[key]=dict(r); merged[alt]=merged[key]; order.append(key); continue
    base=merged[k]
    for f,v in r.items():
        if v and not base.get(f): base[f]=v
        elif f in("주요취급품목","선정이유") and v and v not in base[f] and len(base[f])<120: base[f]=base[f]+" / "+v
    # 우선순위: 중국 자체 소싱망 강함이면 하, 아니면 더 높은 쪽
    if "강함" in base.get("중국자체소싱망","") or "강함" in r.get("중국자체소싱망",""): base["협업가능성"]="하"
    else:
        pr={"상":0,"중":1,"하":2}
        if pr.get(r.get("협업가능성"),1)<pr.get(base.get("협업가능성"),1): base["협업가능성"]=r["협업가능성"]
    if r.get("카테고리") and r["카테고리"] not in base["카테고리"]: base["카테고리"]=base["카테고리"]+" · "+r["카테고리"]
rows=[merged[k] for k in order]
# 정리: 형식 의심 번호, 개인 이름형 메일 제외 (명시적 목록)
BAD_PHONE={"015-8504-6667"}
BAD_MAIL={"yuu0421@saengong.com","ldw9783@bmsmile.com"}
for r in rows:
    if r.get("대표전화") in BAD_PHONE:
        r["문의채널"]=(r.get("문의채널","")+" / 전화번호 형식 의심("+r["대표전화"]+"), 홈페이지에서 확인 필요").strip(" /"); r["대표전화"]=""
    if r.get("이메일") in BAD_MAIL:
        r["문의채널"]=(r.get("문의채널","")+" / 개인 메일로 보여 제외, 대표 메일 확인 필요").strip(" /"); r["이메일"]=""
prio = {"상":0,"중":1,"하":2}
rows.sort(key=lambda r:(prio.get(r.get("협업가능성","중"),1), r.get("카테고리",""), r.get("업체명","")))
print("total", len(rows))

FONT = "Arial"
hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
body_font = Font(name=FONT, size=10)
input_fill = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
prio_fill = {"상":PatternFill("solid", fgColor="C6EFCE"), "중":PatternFill("solid", fgColor="FFEB9C"), "하":PatternFill("solid", fgColor="F4CCCC")}

wb = Workbook()
# ---------- 사용법 ----------
ws0 = wb.active; ws0.title = "사용법"
guide = [
 ["올투게더나우 B2B 영업 리스트 (한국 영업 × 중국 소싱 파트너 모델)"],
 [""],
 ["시트 구성"],
 ["영업리스트", "컨택 대상 업체 전체. 노란색 칸(1차연락일~메모)만 직접 채우면 됩니다."],
 ["요약", "카테고리·우선순위별 업체 수와 연락 진행 현황이 자동 집계됩니다."],
 ["영업멘트", "전화·이메일 첫 컨택용 스크립트 초안."],
 [""],
 ["우선순위 뜻"],
 ["상", "중국 소싱 조직이 약하거나 없고, 굿즈/OEM 물량이 꾸준한 곳. 먼저 연락."],
 ["중", "가능성은 있으나 규모·소싱 방식이 불명확. 2차 연락."],
 ["하", "중국 자체 공장·지사를 이미 보유(경쟁사). 정보 수집용, 컨택 후순위."],
 [""],
 ["주의사항"],
 ["1", "연락처는 각 회사가 홈페이지에 영업 목적으로 공개한 대표번호·대표메일·문의채널만 수집했습니다."],
 ["2", "빈칸은 홈페이지에서 확인이 안 된 항목입니다. 전화 걸 때 직접 물어보고 채우세요."],
 ["3", "연락 전 출처URL을 한 번 열어 현재도 영업 중인지, 번호가 바뀌지 않았는지 확인하세요."],
 ["4", "1차연락결과는 드롭다운(통화/메일발송/관심/견적요청/거절/부재/재연락)에서 고르면 요약 시트에 자동 집계됩니다."],
 [""],
 ["조사일", "2026-09-09 (웹 공개정보 기준)"],
]
for r in guide: ws0.append(r)
ws0["A1"].font = Font(name=FONT, bold=True, size=14)
for c in ("A3","A8","A13"): ws0[c].font = Font(name=FONT, bold=True, size=11)
for row in ws0.iter_rows():
    for c in row:
        if not c.font.bold: c.font = body_font
ws0.column_dimensions["A"].width = 16; ws0.column_dimensions["B"].width = 100

# ---------- 영업리스트 ----------
ws = wb.create_sheet("영업리스트")
cols = ["번호","우선순위","업체명","카테고리","홈페이지","대표전화","이메일","문의채널","소재지","주요취급품목","중국자체소싱망","선정이유","출처URL",
        "1차연락일","연락방법","담당부서/담당자","1차연락결과","후속조치","다음연락일","메모"]
widths = [5,7,22,22,30,16,28,30,12,40,26,50,30,11,10,16,12,30,11,30]
ws.append(cols)
for i,(c,w) in enumerate(zip(cols,widths),1):
    cell = ws.cell(row=1,column=i); cell.font=hdr_font; cell.fill=hdr_fill
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=border
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30
INPUT_START = cols.index("1차연락일")+1
for n,r in enumerate(rows,1):
    vals = [n, r.get("협업가능성",""), r.get("업체명",""), r.get("카테고리",""), r.get("홈페이지",""), r.get("대표전화",""),
            r.get("이메일",""), r.get("문의채널",""), r.get("소재지",""), r.get("주요취급품목",""), r.get("중국자체소싱망",""),
            r.get("선정이유",""), r.get("출처URL",""), "","","","","","",""]
    ws.append(vals)
    rr = n+1
    for ci in range(1,len(cols)+1):
        c = ws.cell(row=rr,column=ci); c.font=body_font; c.border=border
        c.alignment=Alignment(vertical="top",wrap_text=True)
        if ci>=INPUT_START: c.fill=input_fill
    p = r.get("협업가능성","")
    if p in prio_fill: ws.cell(row=rr,column=2).fill = prio_fill[p]; ws.cell(row=rr,column=2).alignment=Alignment(horizontal="center",vertical="top")
    for ci in (5,13):
        c = ws.cell(row=rr,column=ci)
        if c.value and str(c.value).startswith("http"):
            c.hyperlink = c.value; c.font = Font(name=FONT,size=10,color="0563C1",underline="single")
last = len(rows)+1
dv = DataValidation(type="list", formula1='"통화,메일발송,관심,견적요청,거절,부재,재연락"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"Q2:Q{max(last,2)}")
dv2 = DataValidation(type="list", formula1='"전화,이메일,홈페이지문의,카톡채널,방문"', allow_blank=True)
ws.add_data_validation(dv2); dv2.add(f"O2:O{max(last,2)}")
ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
ws.cell(row=1,column=INPUT_START).comment = Comment("노란색 칸은 대표님이 직접 입력하는 칸입니다.", "Claude")

# ---------- 요약 ----------
sm = wb.create_sheet("요약")
sm["A1"] = "카테고리별 업체 수"; sm["A1"].font = Font(name=FONT,bold=True,size=12)
sm.append(["카테고리","전체","상","중","하","연락완료","관심/견적요청"])
cats = sorted(set(r.get("카테고리","") for r in rows))
L = "영업리스트"
N = len(rows)+300  # 추가 입력 여유분
for i,cat in enumerate(cats, start=3):
    sm.cell(row=i,column=1,value=cat)
    sm.cell(row=i,column=2,value=f"=COUNTIF({L}!$D$2:$D${N},A{i})")
    for j,p in enumerate(["상","중","하"], start=3):
        sm.cell(row=i,column=j,value=f'=COUNTIFS({L}!$D$2:$D${N},A{i},{L}!$B$2:$B${N},"{p}")')
    sm.cell(row=i,column=6,value=f'=COUNTIFS({L}!$D$2:$D${N},A{i},{L}!$Q$2:$Q${N},"<>")')
    sm.cell(row=i,column=7,value=f'=COUNTIFS({L}!$D$2:$D${N},A{i},{L}!$Q$2:$Q${N},"관심")+COUNTIFS({L}!$D$2:$D${N},A{i},{L}!$Q$2:$Q${N},"견적요청")')
tr = 3+len(cats)
sm.cell(row=tr,column=1,value="합계").font=Font(name=FONT,bold=True)
for j in range(2,8):
    col = get_column_letter(j)
    sm.cell(row=tr,column=j,value=f"=SUM({col}3:{col}{tr-1})").font=Font(name=FONT,bold=True)
for row in sm.iter_rows(min_row=2,max_row=tr):
    for c in row:
        c.border=border
        if not c.font.bold: c.font=body_font
for j in range(1,8):
    c=sm.cell(row=2,column=j); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(horizontal="center")
sm.column_dimensions["A"].width=40
for j in range(2,8): sm.column_dimensions[get_column_letter(j)].width=13
sm.cell(row=tr+2,column=1,value="연락결과 현황").font=Font(name=FONT,bold=True,size=12)
sm.cell(row=tr+3,column=1,value="결과"); sm.cell(row=tr+3,column=2,value="건수")
for j in (1,2):
    c=sm.cell(row=tr+3,column=j); c.font=hdr_font; c.fill=hdr_fill
for k,res in enumerate(["통화","메일발송","관심","견적요청","거절","부재","재연락"]):
    rr=tr+4+k
    sm.cell(row=rr,column=1,value=res).font=body_font
    sm.cell(row=rr,column=2,value=f"=COUNTIF({L}!$Q$2:$Q${N},A{rr})").font=body_font

# ---------- 영업멘트 ----------
sc = wb.create_sheet("영업멘트")
sc.column_dimensions["A"].width=14; sc.column_dimensions["B"].width=110
lines = [
 ["전화 첫 멘트", "안녕하세요, 올투게더나우 대표 ○○○입니다. 저희는 무릎·허리 보호대를 중국과 한국에서 직접 기획·제조해서 네이버·쿠팡에 판매하는 회사이고, 수출입 사업자를 갖고 있습니다. 귀사에서 굿즈나 판촉물 제작하실 때 중국 생산이 들어가는 품목이 있으면, 공장 발굴부터 샘플, 가격협상, 생산, 검수, 국내 입고까지 저희가 한 번에 맡아드릴 수 있어서 연락드렸습니다. 구매 담당 부서로 연결 가능할까요?"],
 ["핵심 차별점", "1) 단순 구매대행이 아니라 중국 현지 제조 파트너가 직접 공장을 찾고 생산을 관리합니다. 2) 저희가 실제로 자사 제품을 중국에서 만들어 팔고 있어서 품질·납기 사고를 직접 겪어본 경험이 있습니다. 3) 수출입 사업자로 세금계산서 발행, 통관, KC 인증 안내까지 국내에서 처리합니다."],
 ["이메일 제목", "[제안] 굿즈·판촉물 중국 OEM 생산 대행 (공장 발굴~검수~입고 원스톱) - 올투게더나우"],
 ["이메일 본문", "안녕하세요, 올투게더나우 대표 ○○○입니다.\n\n저희는 신체 보호대(무릎·팔꿈치·손목·허리)를 중국·한국에서 기획·제조하여 네이버 스마트스토어와 쿠팡에서 판매하고 있으며, 수출입 사업자를 보유하고 있습니다.\n\n귀사에서 고객사 굿즈·판촉물·기념품 제작 시 중국 생산이 필요한 품목이 있으시면, 아래 과정을 저희가 원스톱으로 진행해 드립니다.\n\n1. 필요한 제품·사양·수량·예산만 알려주시면 중국 현지 파트너가 제조공장을 발굴\n2. 샘플 제작 및 국내 발송\n3. 가격 협상 및 견적서 제출\n4. 생산 진행 및 출하 전 검수 (사진·영상 보고)\n5. 통관 후 국내 지정 장소 납품, 세금계산서 발행\n\n부담 없이 진행 중인 건 하나로 견적 비교부터 해보셔도 좋습니다.\n연락 주시면 바로 찾아뵙고 설명드리겠습니다.\n\n올투게더나우 대표 ○○○ / 전화 000-0000-0000 / 이메일 ○○○@○○○"],
 ["예상 질문 1", "Q. 이미 거래하는 중국 공장이 있는데요? → A. 그러면 견적 비교용으로만 한 건 받아보세요. 같은 사양으로 저희 견적이 더 낮거나 납기가 빠르면 그때 바꾸시면 됩니다."],
 ["예상 질문 2", "Q. 소량도 되나요? → A. 품목별로 다르지만 300~500개부터 가능한 공장을 찾아드립니다. 우선 사양 주시면 MOQ부터 확인해 드릴게요."],
 ["예상 질문 3", "Q. 불량 나면 어떻게 하나요? → A. 출하 전 검수 사진·영상을 보내드리고, 검수 통과 후에만 출고합니다. 계약서에 불량 시 재생산 조건을 명시합니다."],
 ["후속 조치", "통화 후 24시간 안에 이메일로 회사소개(보호대 제조 사례 사진, 공장 사진) + 견적 요청 양식(품목/사양/수량/희망단가/납기) 발송. 1주 뒤 재연락."],
]
for r in lines: sc.append(r)
for row in sc.iter_rows():
    row[0].font=Font(name=FONT,bold=True); row[0].alignment=Alignment(vertical="top")
    row[1].font=body_font; row[1].alignment=Alignment(wrap_text=True,vertical="top")

os.makedirs(OUT_DIR, exist_ok=True)
xlsx = f"{OUT_DIR}/B2B_영업리스트.xlsx"
wb.calculation.fullCalcOnLoad = True
wb.save(xlsx); print("saved", xlsx)

# CSV too
with open(f"{OUT_DIR}/B2B_영업리스트.csv","w",newline="",encoding="utf-8-sig") as fp:
    w = csv.writer(fp); w.writerow(cols[:13])
    for n,r in enumerate(rows,1):
        w.writerow([n, r.get("협업가능성",""), r.get("업체명",""), r.get("카테고리",""), r.get("홈페이지",""), r.get("대표전화",""),
            r.get("이메일",""), r.get("문의채널",""), r.get("소재지",""), r.get("주요취급품목",""), r.get("중국자체소싱망",""),
            r.get("선정이유",""), r.get("출처URL","")])
json.dump(rows, open(f"{OUT_DIR}/B2B_영업리스트.json","w",encoding="utf-8"), ensure_ascii=False, indent=1)
